#!/usr/bin/env python3
"""
Lottery predictor (for fun) - Powerball and Mega Millions
=========================================================
One script that:
  1. UPDATES its data files from the official NY Open Data lottery feeds
     (data.ny.gov), appending only draws it doesn't have yet.
  2. Builds 5 guesses per game for the next draw using 5 strategies.
  3. LOGS every guess before the draw, and SCORES them automatically on the
     next run once the real results arrive.
  4. Rebuilds an Excel workbook (draws, frequency, next picks, scorecard).
  5. Can backtest any strategy on past draws (--backtest).

Usage
-----
  python lotto_predictor.py                 # update data, predict next draws, write workbook
  python lotto_predictor.py --no-update     # skip the download (offline)
  python lotto_predictor.py --game powerball
  python lotto_predictor.py --date 2026-03-23   # what would it have guessed for that date?
  python lotto_predictor.py --backtest      # walk-forward test of every strategy
  python lotto_predictor.py --days 365      # history window (default 365 days)

Run it regularly (e.g. after each draw). Examples:
  Linux/macOS cron (runs 7 AM daily):   0 7 * * *  cd /path/to/lotto_kit && python3 lotto_predictor.py
  Windows Task Scheduler: action = python.exe, arguments = lotto_predictor.py, start in = this folder.

Requires Python 3.9+. The Excel workbook needs:  pip install openpyxl
This is a for-fun math experiment. Every combination has identical odds; no formula
predicts lottery draws.
"""
import argparse, csv, itertools, json, math, os, random, statistics, sys, urllib.parse, urllib.request
import datetime as dt
from collections import Counter

try:
    from zoneinfo import ZoneInfo
except Exception:          # very old Python / missing tzdata
    ZoneInfo = None

HERE = os.path.dirname(os.path.abspath(__file__))

# --------------------------------------------------------------------------- games
GAMES = {
    "powerball": dict(
        label="Powerball", dataset="d6yy-54nr", whites=69, bonus=26, bonus_name="Powerball",
        weekdays={0, 2, 5}, draw_time="10:59 PM ET", matrix_start=dt.date(2015, 10, 7)),
    "megamillions": dict(
        label="Mega Millions", dataset="5xaw-6ayf", whites=70, bonus=24, bonus_name="Mega Ball",
        weekdays={1, 4}, draw_time="11:00 PM ET",
        # Mega Millions changed its number matrix in April 2025 (bonus ball 1-25 -> 1-24).
        # Draws before this date are ignored so every game uses the current format.
        matrix_start=dt.date(2025, 4, 8)),
}
SALES_CUTOFF_HOUR_ET = 22   # after this hour we target the following draw
FEED = "https://data.ny.gov/resource/{ds}.json"

# --------------------------------------------------------------------------- data files
def data_path(game, data_dir):
    return os.path.join(data_dir, f"{game}_draws.csv")

def load_draws(game, data_dir):
    """Returns list of (date, [5 whites sorted], bonus), newest first."""
    p = data_path(game, data_dir)
    out = []
    if os.path.exists(p):
        with open(p, newline="") as f:
            for r in csv.DictReader(f):
                out.append((dt.date.fromisoformat(r["date"]),
                            sorted(int(r[f"n{i}"]) for i in range(1, 6)), int(r["bonus"])))
    return sorted(out, key=lambda x: x[0], reverse=True)

def save_draws(game, data_dir, draws):
    os.makedirs(data_dir, exist_ok=True)
    with open(data_path(game, data_dir), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["date", "n1", "n2", "n3", "n4", "n5", "bonus"])
        for d, whites, b in sorted(draws, key=lambda x: x[0]):
            w.writerow([d.isoformat(), *whites, b])

def fetch_json(url):
    req = urllib.request.Request(url, headers={"User-Agent": "lotto-predictor/1.0"})
    with urllib.request.urlopen(req, timeout=25) as resp:
        return json.loads(resp.read().decode("utf-8"))

def parse_feed_row(game, row):
    """Convert one feed record to (date, whites, bonus)."""
    d = dt.date.fromisoformat(row["draw_date"][:10])
    nums = [int(x) for x in row["winning_numbers"].split()]
    if game == "powerball":
        whites, bonus = nums[:5], nums[5]
    else:
        whites, bonus = nums[:5], int(row["mega_ball"])
    return d, sorted(whites), bonus

def update_data(game, data_dir):
    """Download draws newer than what we have. Returns (new_count, message)."""
    cfg = GAMES[game]
    have = load_draws(game, data_dir)
    since = have[0][0] if have else cfg["matrix_start"] - dt.timedelta(days=1)
    q = urllib.parse.urlencode({"$where": f"draw_date > '{since.isoformat()}'",
                                "$order": "draw_date", "$limit": "2000"})
    try:
        rows = fetch_json(FEED.format(ds=cfg["dataset"]) + "?" + q)
    except Exception as e:
        return 0, f"could not reach the lottery feed ({type(e).__name__}: {e}); using saved data"
    known = {d for d, _, _ in have}
    new = []
    for r in rows:
        try:
            d, w, b = parse_feed_row(game, r)
        except Exception:
            continue
        if d not in known and d >= cfg["matrix_start"]:
            new.append((d, w, b)); known.add(d)
    if new:
        save_draws(game, data_dir, have + new)
    return len(new), (f"added {len(new)} new draw(s), latest {max(d for d, _, _ in new)}" if new
                      else "already up to date")

def ingest_file(game, path, data_dir):
    """Add draws from a text file. Accepts the feed's JSON, or lines like
    'YYYY-MM-DD 03 34 37 42 46 20' (five white balls then the bonus ball). Returns count added."""
    cfg = GAMES[game]
    text = open(path, encoding="utf-8").read()
    parsed = []
    if "[" in text and "{" in text:
        try:
            for r in json.loads(text[text.index("["): text.rindex("]") + 1]):
                try: parsed.append(parse_feed_row(game, r))
                except Exception: pass
        except Exception:
            pass
    if not parsed:
        for line in text.splitlines():
            t = line.replace("|", " ").replace(",", " ").split()
            if len(t) == 7:
                try:
                    d = dt.date.fromisoformat(t[0]); nums = [int(x) for x in t[1:]]
                except Exception:
                    continue
                parsed.append((d, sorted(nums[:5]), nums[5]))
    have = load_draws(game, data_dir); known = {d for d, _, _ in have}; new = []
    for d, w, b in parsed:
        ok_ = len(set(w)) == 5 and all(1 <= x <= cfg["whites"] for x in w) and 1 <= b <= cfg["bonus"]
        if ok_ and d not in known and d >= cfg["matrix_start"]:
            new.append((d, w, b)); known.add(d)
    if new: save_draws(game, data_dir, have + new)
    return len(new)

# --------------------------------------------------------------------------- dates
def now_et():
    if ZoneInfo:
        try:
            return dt.datetime.now(ZoneInfo("America/New_York"))
        except Exception:
            pass
    return dt.datetime.now()

def next_draw_date(game, start_date):
    d = start_date
    while d.weekday() not in GAMES[game]["weekdays"]:
        d += dt.timedelta(days=1)
    return d

def default_start():
    n = now_et()
    return n.date() + dt.timedelta(days=1) if n.hour >= SALES_CUTOFF_HOUR_ET else n.date()

# --------------------------------------------------------------------------- scoring helpers
def zscores(d):
    vals = list(d.values())
    mu, sd = statistics.mean(vals), statistics.pstdev(vals) or 1
    return {k: (v - mu) / sd for k, v in d.items()}

def score_components(draws_numbers, pool_size, half=20):
    """Standardized frequency / due / hot-streak signals. draws_numbers is newest first."""
    N, k = len(draws_numbers), len(draws_numbers[0])
    freq = {n: 0 for n in range(1, pool_size + 1)}
    due = {n: N for n in freq}
    hot = {n: 0.0 for n in freq}
    seen = set()
    for i, nums in enumerate(draws_numbers):
        w = 0.5 ** (i / half)                     # hot-streak half-life in draws
        for n in nums:
            freq[n] += 1; hot[n] += w
            if n not in seen: due[n] = i; seen.add(n)
    gap = pool_size / k
    due = {n: due[n] / gap for n in due}
    return zscores(freq), zscores(due), zscores(hot)

def score_pool(draws_numbers, pool_size, wf=0.45, wd=0.25, wh=0.30, half=20):
    """Blend of the three signals. Defaults are the original formula."""
    zf, zd, zh = score_components(draws_numbers, pool_size, half)
    return {n: wf * zf[n] + wd * zd[n] + wh * zh[n] for n in zf}

def ok(combo, lo, hi, whites):
    """Shape filters used by every strategy."""
    c = sorted(combo)
    if not (lo <= sum(c) <= hi): return False                      # typical sum band
    if sum(x % 2 for x in c) not in (2, 3): return False           # 2-3 odd numbers
    if sum(x <= whites // 2 for x in c) not in (2, 3): return False   # 2-3 low numbers
    if sum(1 for a, b in zip(c, c[1:]) if b - a == 1) > 1: return False  # <=1 consecutive pair
    return True

def sum_band(H):
    s = sorted(sum(w) for w, _ in H)
    return s[int(.25 * (len(s) - 1))], s[int(.75 * (len(s) - 1))]

def best_combo(cands, value, lo, hi, whites):
    best, bv = None, -1e18
    for combo in itertools.combinations(cands, 5):
        if not ok(combo, lo, hi, whites): continue
        v = value(combo)
        if v > bv: best, bv = sorted(combo), v
    return best

def ew_scores(H, pool, half, use_bonus):
    sc = {n: 0.0 for n in range(1, pool + 1)}
    for i, (w, b) in enumerate(H):
        wt = 0.5 ** (i / half)
        for n in ([b] if use_bonus else w): sc[n] += wt
    return sc

def gaps(H, pool, use_bonus):
    g, seen = {n: len(H) for n in range(1, pool + 1)}, set()
    for i, (w, b) in enumerate(H):
        for n in ([b] if use_bonus else w):
            if n not in seen: g[n] = i; seen.add(n)
    return g

# --------------------------------------------------------------------------- strategies
# Each takes H = [(whites, bonus), ...] newest first, and the game config.
def s_formula(H, cfg):
    W, S = [w for w, _ in H], [[b] for _, b in H]
    lo, hi = sum_band(H)
    ws = score_pool(W, cfg["whites"]); ss = score_pool(S, cfg["bonus"])
    ranked = sorted(ws, key=ws.get, reverse=True)[:18]
    combo = best_combo(ranked, lambda c: sum(ws[x] for x in c), lo, hi, cfg["whites"])
    return combo, max(ss, key=ss.get)

def s_balanced(H, cfg):
    lo, hi = sum_band(H); nw = cfg["whites"]
    freq = Counter(x for w, _ in H for x in w)
    cands = sorted(range(1, nw + 1), key=lambda n: -freq.get(n, 0))[:24]
    med = statistics.median(sum(w) for w, _ in H)
    def val(c):
        dec = len({(x - 1) // 10 for x in c})
        return 3 * dec - abs(sum(c) - med) / 10 + 0.1 * sum(freq.get(x, 0) for x in c)
    combo = best_combo(cands, val, lo, hi, nw)
    pbf = Counter(b for _, b in H)
    mid = range(round(cfg["bonus"] * 0.3), round(cfg["bonus"] * 0.75))
    return combo, max(mid, key=lambda n: pbf.get(n, 0))

def s_overdue(H, cfg):
    lo, hi = sum_band(H)
    g = gaps(H, cfg["whites"], False)
    cands = sorted(g, key=lambda n: -g[n])[:16]
    combo = best_combo(cands, lambda c: sum(g[x] for x in c), lo, hi, cfg["whites"])
    gp = gaps(H, cfg["bonus"], True)
    return combo, max(gp, key=gp.get)

def s_hot(H, cfg):
    lo, hi = sum_band(H)
    sc = ew_scores(H, cfg["whites"], 8, False)
    cands = sorted(sc, key=lambda n: -sc[n])[:18]
    combo = best_combo(cands, lambda c: sum(sc[x] for x in c), lo, hi, cfg["whites"])
    ps = ew_scores(H, cfg["bonus"], 8, True)
    return combo, max(ps, key=ps.get)

def s_pairs(H, cfg):
    lo, hi = sum_band(H); nw = cfg["whites"]
    pairs = Counter(p for w, _ in H for p in itertools.combinations(w, 2))
    freq = Counter(x for w, _ in H for x in w)
    cands = sorted(range(1, nw + 1), key=lambda n: -freq.get(n, 0))[:22]
    val = lambda c: (sum(pairs.get(p, 0) for p in itertools.combinations(sorted(c), 2))
                     + 0.05 * sum(freq.get(x, 0) for x in c))
    combo = best_combo(cands, val, lo, hi, nw)
    pw = Counter()
    for w, b in H:
        ov = len(set(w) & set(combo)); pw[b] += ov * ov
    return combo, max(pw, key=lambda n: (pw[n], n))

STRATEGIES = [
    ("Formula (frequency + due + hot)", s_formula),
    ("Balanced profile", s_balanced),
    ("Overdue gaps", s_overdue),
    ("Hot streak (fast decay)", s_hot),
    ("Pair affinity", s_pairs),
]
STRATEGY_HELP = {
    "Formula (frequency + due + hot)": "45% frequency + 25% due + 30% hot-streak score; best combo passing shape filters.",
    "Balanced profile": "Frequent numbers; most decade variety, sum near the median.",
    "Overdue gaps": "Numbers unseen the longest; bonus ball unseen the longest.",
    "Hot streak (fast decay)": "Recent draws weighted heavily (half-life 8 draws).",
    "Pair affinity": "Combo whose numbers have appeared together most often.",
    "Adaptive (self-tuning)": "Re-tunes the formula's weights and hot-streak speed on recent draws (walk-forward); only adopts a change if it beats the default by a wide statistical margin.",
}
ADAPTIVE_NAME = "Adaptive (self-tuning)"
ALL_NAMES = [n for n, _ in STRATEGIES] + [ADAPTIVE_NAME]

# --------------------------------------------------------------------------- history window
def history_for(draws, target, days, cfg):
    start = max(target - dt.timedelta(days=days), cfg["matrix_start"])
    return [(w, b) for d, w, b in draws if start <= d < target], start

# --------------------------------------------------------------------------- self-tuning
# The script "learns" from each new winning draw like this: for the last `window` draws it
# replays history (using only draws BEFORE each one), asks how highly the formula ranked the
# numbers that actually won, and searches a small grid of weights / hot-streak speeds for the
# setting that ranked the winners best. The metric is the average rank of the winning numbers
# as a share of the pool (0 = always ranked first, 0.50 = no better than random).
# Because lottery draws are random, most grid settings will only look better by luck. So a
# change is adopted only when it beats the default by z >= Z_ADOPT (about 99.4% confidence
# for a single test, still strict after trying ~38 settings). Otherwise the default is kept.
GRID_WEIGHTS = [(1, 0, 0), (0, 1, 0), (0, 0, 1), (.45, .25, .30), (.5, .25, .25), (.25, .5, .25),
                (.25, .25, .5), (.34, .33, .33), (.6, .2, .2), (.2, .2, .6), (0, .5, .5),
                (.5, 0, .5), (.5, .5, 0)]
GRID_HALF = [8, 20, 40]
DEFAULT_PARAMS = (0.45, 0.25, 0.30, 20)
Z_ADOPT = 2.5

def rank_pct(scores, winners):
    order = sorted(scores, key=scores.get, reverse=True)
    pos = {n: i for i, n in enumerate(order)}
    return statistics.mean(pos[w] for w in winners) / (len(order) - 1)

def build_tests(draws, cfg, target, days, window):
    chrono = sorted([x for x in draws if x[0] < target], key=lambda x: x[0])
    tests = []
    for d, w, b in chrono[-window:]:
        H, _ = history_for(draws, d, days, cfg)
        if len(H) >= 30: tests.append((H, w, b))
    return tests

def eval_params(params, tests, cfg, bonus):
    wf, wd, wh, half = params
    out = []
    for H, w, b in tests:
        if bonus:
            sc = score_pool([[bb] for _, bb in H], cfg["bonus"], wf, wd, wh, half)
            out.append(rank_pct(sc, [b]))
        else:
            sc = score_pool([ww for ww, _ in H], cfg["whites"], wf, wd, wh, half)
            out.append(rank_pct(sc, w))
    return out

def tune(draws, cfg, target, days, window=60):
    tests = build_tests(draws, cfg, target, days, window)
    res = {}
    for key, bonus in (("whites", False), ("bonus", True)):
        base = eval_params(DEFAULT_PARAMS, tests, cfg, bonus)
        best, best_z, best_vals = DEFAULT_PARAMS, 0.0, base
        for wts in GRID_WEIGHTS:
            for half in GRID_HALF:
                p = (*wts, half)
                if p == DEFAULT_PARAMS: continue
                vals = eval_params(p, tests, cfg, bonus)
                diffs = [a - b for a, b in zip(base, vals)]      # >0 means candidate ranks winners better
                sd = statistics.pstdev(diffs) or 1e-9
                z = statistics.mean(diffs) / (sd / math.sqrt(len(diffs)))
                if z > best_z: best, best_z, best_vals = p, z, vals
        adopted = best_z >= Z_ADOPT
        res[key] = dict(params=best if adopted else DEFAULT_PARAMS, adopted=adopted, z=best_z,
                        best_candidate=best, default_metric=statistics.mean(base),
                        chosen_metric=statistics.mean(best_vals if adopted else base),
                        candidate_metric=statistics.mean(best_vals), n=len(tests))
    return res

def make_adaptive(tuned):
    def s_adaptive(H, cfg):
        lo, hi = sum_band(H)
        ws = score_pool([w for w, _ in H], cfg["whites"], *tuned["whites"]["params"])
        ss = score_pool([[b] for _, b in H], cfg["bonus"], *tuned["bonus"]["params"])
        ranked = sorted(ws, key=ws.get, reverse=True)[:18]
        combo = best_combo(ranked, lambda c: sum(ws[x] for x in c), lo, hi, cfg["whites"])
        return combo, max(ss, key=ss.get)
    return s_adaptive

def learn_note(tuned):
    w, b = tuned["whites"], tuned["bonus"]
    def one(x, label):
        if x["adopted"]:
            wf, wd, wh, half = x["params"]
            return (f"{label}: adopted weights freq {wf:.2f}/due {wd:.2f}/hot {wh:.2f}, half-life {half} "
                    f"(winners ranked in top {100 * x['chosen_metric']:.1f}% vs {100 * x['default_metric']:.1f}% default)")
        return (f"{label}: no setting beat the default by a reliable margin (best looked like z={x['z']:.1f}, needs {Z_ADOPT}); "
                f"winners ranked in top {100 * x['default_metric']:.1f}% on average; random = 50%")
    return one(w, "white balls") + " | " + one(b, "bonus ball")

def audit(game, draws, days, n_test, window):
    """Honest out-of-sample test of self-tuning: tune using only data before each test draw."""
    cfg = GAMES[game]
    chrono = sorted(draws, key=lambda x: x[0])[-n_test:]
    adap_w, base_w, adap_rank, base_rank, adopted = [], [], [], [], 0
    for d, w, b in chrono:
        tuned = tune(draws, cfg, d, days, window)
        H, _ = history_for(draws, d, days, cfg)
        c, _ = make_adaptive(tuned)(H, cfg); c0, _ = s_formula(H, cfg)
        adap_w.append(len(set(w) & set(c))); base_w.append(len(set(w) & set(c0)))
        sa = score_pool([x for x, _ in H], cfg["whites"], *tuned["whites"]["params"])
        sb = score_pool([x for x, _ in H], cfg["whites"])
        adap_rank.append(rank_pct(sa, w)); base_rank.append(rank_pct(sb, w))
        adopted += tuned["whites"]["adopted"]
    n = len(chrono)
    print(f"\n{cfg['label']} out-of-sample audit over the last {n} draws (tuned using only earlier draws):")
    print(f"  Self-tuning changed the default weights in {adopted} of {n} draws")
    print(f"  Winning white numbers ranked in top: adaptive {100 * statistics.mean(adap_rank):.1f}% | default formula {100 * statistics.mean(base_rank):.1f}% | random 50.0%")
    print(f"  Avg white matches per guess:         adaptive {statistics.mean(adap_w):.2f} | default formula {statistics.mean(base_w):.2f} | random {25 / cfg['whites']:.2f}")

def review(game, draws, days, k):
    """Where did the model go wrong on the last k draws, and which signal was closest?"""
    cfg = GAMES[game]
    chrono = sorted(draws, key=lambda x: x[0])[-k:]
    agg = {"frequency": [], "due": [], "hot streak": [], "blended formula": []}
    print(f"\n{cfg['label']} - review of the last {len(chrono)} draws (model sees only earlier draws):")
    for d, w, b in chrono:
        H, _ = history_for(draws, d, days, cfg)
        if len(H) < 30: continue
        zf, zd, zh = score_components([x for x, _ in H], cfg["whites"])
        blend = score_pool([x for x, _ in H], cfg["whites"])
        order = sorted(blend, key=blend.get, reverse=True); pos = {n: i + 1 for i, n in enumerate(order)}
        print(f"\n  {d:%a %Y-%m-%d}: winning {' '.join(f'{x:02d}' for x in w)} + {b:02d}")
        print("    Where the formula ranked each winner (1 = top pick of "
              f"{cfg['whites']}): " + ", ".join(f"{x:02d}->#{pos[x]}" for x in w))
        for label, sc in (("frequency", zf), ("due", zd), ("hot streak", zh), ("blended formula", blend)):
            agg[label].append(rank_pct(sc, w))
        for name, f in STRATEGIES:
            c, pb = f(H, cfg)
            print(f"    {name:34s} guessed {' '.join(f'{x:02d}' for x in c)} + {pb:02d}: "
                  f"{len(set(c) & set(w))} white match(es), bonus {'hit' if pb == b else 'miss'}")
    if agg["blended formula"]:
        print("\n  Average rank of winning numbers by signal (lower is closer; random = 50%):")
        for label, vals in sorted(agg.items(), key=lambda kv: statistics.mean(kv[1])):
            print(f"    {label:16s} top {100 * statistics.mean(vals):.1f}%")

# --------------------------------------------------------------------------- backtest
def backtest(game, draws, days, start_after=40):
    cfg = GAMES[game]
    chrono = sorted(draws, key=lambda x: x[0])
    random.seed(11)
    res = {n: dict(w=[], b=[]) for n, _ in STRATEGIES}
    rw, rb = [], []
    for i, (d, w, b) in enumerate(chrono):
        H, _ = history_for(draws, d, days, cfg)
        if len(H) < start_after: continue
        for name, f in STRATEGIES:
            c, pb = f(H, cfg)
            res[name]["w"].append(len(set(w) & set(c))); res[name]["b"].append(pb == b)
        c = random.sample(range(1, cfg["whites"] + 1), 5)
        rw.append(len(set(w) & set(c))); rb.append(random.randint(1, cfg["bonus"]) == b)
    return res, rw, rb

# --------------------------------------------------------------------------- prediction log
LOG_FIELDS = ["game", "target_date", "strategy", "whites", "bonus", "generated_at"]

def load_log(path):
    if not os.path.exists(path): return []
    with open(path, newline="") as f: return list(csv.DictReader(f))

def append_log(path, rows):
    new = not os.path.exists(path)
    with open(path, "a", newline="") as f:
        w = csv.DictWriter(f, fieldnames=LOG_FIELDS)
        if new: w.writeheader()
        w.writerows(rows)

def scored_log(log, all_draws):
    """Attach actual results and match counts to each logged prediction."""
    out = []
    for r in log:
        g = r["game"]; td = dt.date.fromisoformat(r["target_date"])
        actual = next(((w, b) for d, w, b in all_draws.get(g, []) if d == td), None)
        pw = [int(x) for x in r["whites"].split()]; pb = int(r["bonus"])
        row = dict(r, target=td, pw=pw, pb=pb, actual=actual)
        if actual:
            row["wm"] = len(set(pw) & set(actual[0])); row["bm"] = pb == actual[1]
        out.append(row)
    return sorted(out, key=lambda x: (x["target"], x["game"]))

# --------------------------------------------------------------------------- workbook
def write_workbook(path, all_draws, preds, scored, days, asof_start):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("  (openpyxl not installed - skipping Excel file. Run: pip install openpyxl)")
        return False
    F = "Arial"
    hf = Font(name=F, bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="1F3864")
    body = Font(name=F); bold = Font(name=F, bold=True); blue = Font(name=F, bold=True, color="0000FF")
    thin = Side(style="thin", color="BBBBBB"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bfill = {"powerball": "F8CBAD", "megamillions": "FFE699"}
    wb = Workbook(); wb.remove(wb.active)

    # ---- Next Picks
    ws = wb.create_sheet("Next Picks")
    ws["A1"] = "Guesses for the next draws"; ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = f"Generated {dt.datetime.now():%Y-%m-%d %H:%M}. For-fun experiment; every combination has identical odds."
    ws["A2"].font = Font(name=F, italic=True)
    r = 4
    for g, p in preds.items():
        cfg = GAMES[g]
        ws.cell(r, 1, f"{cfg['label']}  -  draw on {p['target']:%a %Y-%m-%d} ({cfg['draw_time']})").font = Font(name=F, bold=True, size=12)
        ws.cell(r + 1, 1, f"Based on {p['n']} draws from {p['start']:%Y-%m-%d} to {p['last']:%Y-%m-%d}").font = Font(name=F, italic=True)
        heads = ["Strategy", "Ball 1", "Ball 2", "Ball 3", "Ball 4", "Ball 5", cfg["bonus_name"], "Sum", "How it picks"]
        for c, h in enumerate(heads, start=1):
            cell = ws.cell(r + 2, c, h); cell.font = hf; cell.fill = hfill; cell.alignment = ctr; cell.border = box
        for i, (name, combo, bonus) in enumerate(p["picks"]):
            rr = r + 3 + i
            ws.cell(rr, 1, name)
            for c, v in enumerate(combo, start=2): ws.cell(rr, c, v)
            ws.cell(rr, 7, bonus); ws.cell(rr, 8, f"=SUM(B{rr}:F{rr})"); ws.cell(rr, 9, STRATEGY_HELP[name])
            for c in range(1, 10):
                cell = ws.cell(rr, c); cell.border = box
                cell.font = blue if 2 <= c <= 7 else body
                cell.alignment = ctr if 2 <= c <= 8 else Alignment(vertical="center", wrap_text=(c == 9))
            ws.cell(rr, 7).fill = PatternFill("solid", fgColor=bfill[g])
        if p.get("learn_note"):
            ws.cell(r + 3 + len(p["picks"]), 1, "Learning check: " + p["learn_note"]).font = Font(name=F, italic=True, size=9)
        r += 13
    for col, w in zip("ABCDEFGHI", [34, 8, 8, 8, 8, 8, 12, 8, 70]): ws.column_dimensions[col].width = w

    # ---- Draw sheets
    ranges = {}
    for g in preds:
        cfg = GAMES[g]
        sh = wb.create_sheet(f"{cfg['label']} Draws")
        for c, h in enumerate(["Draw Date", "Day", "Ball 1", "Ball 2", "Ball 3", "Ball 4", "Ball 5", cfg["bonus_name"], "Sum of 5"], start=1):
            cell = sh.cell(1, c, h); cell.font = hf; cell.fill = hfill; cell.alignment = ctr
        rows = all_draws[g]
        for i, (d, w, b) in enumerate(rows, start=2):
            sh.cell(i, 1, d).number_format = "yyyy-mm-dd"
            sh.cell(i, 2, f'=TEXT(A{i},"ddd")')
            for c, v in enumerate(w, start=3): sh.cell(i, c, v)
            sh.cell(i, 8, b); sh.cell(i, 9, f"=SUM(C{i}:G{i})")
            for c in range(1, 10):
                sh.cell(i, c).font = body
                if c > 1: sh.cell(i, c).alignment = Alignment(horizontal="center")
            sh.cell(i, 8).fill = PatternFill("solid", fgColor=bfill[g]); sh.cell(i, 8).font = bold
        sh.freeze_panes = "A2"; sh.auto_filter.ref = f"A1:I{len(rows) + 1}"
        for col, w in zip("ABCDEFGHI", [13, 7, 8, 8, 8, 8, 8, 12, 12]): sh.column_dimensions[col].width = w
        ranges[g] = (sh.title, len(rows) + 1)

    # ---- Frequency (formulas, window starts at the date in B3)
    fs = wb.create_sheet("Frequency")
    fs["A1"] = "Number frequency inside the analysis window"; fs["A1"].font = Font(name=F, bold=True, size=14)
    fs["A3"] = "Window starts (edit to change):"; fs["A3"].font = bold
    fs["D3"] = asof_start; fs["D3"].number_format = "yyyy-mm-dd"; fs["D3"].font = blue
    fs["D3"].fill = PatternFill("solid", fgColor="FFFF00")
    col = 1
    for g in preds:
        cfg = GAMES[g]; sheet, last = ranges[g]; q = f"'{sheet}'"
        for c, h in enumerate([f"{cfg['label']} number", "White ball count", f"{cfg['bonus_name']} count"], start=col):
            cell = fs.cell(5, c, h); cell.font = hf; cell.fill = hfill; cell.alignment = ctr
        for n in range(1, cfg["whites"] + 1):
            rr = 5 + n
            fs.cell(rr, col, n)
            fs.cell(rr, col + 1, f"=SUMPRODUCT(({q}!$A$2:$A${last}>=$D$3)*({q}!$C$2:$G${last}={get_col(col)}{rr}))")
            if n <= cfg["bonus"]:
                fs.cell(rr, col + 2, f"=SUMPRODUCT(({q}!$A$2:$A${last}>=$D$3)*({q}!$H$2:$H${last}={get_col(col)}{rr}))")
            for c in range(col, col + 3): fs.cell(rr, c).font = body; fs.cell(rr, c).alignment = Alignment(horizontal="center")
        for c, w in zip(range(col, col + 3), [16, 16, 18]): fs.column_dimensions[get_col(c)].width = w
        col += 4
    fs.row_dimensions[5].height = 32; fs.freeze_panes = "A6"

    # ---- Scorecard
    sc = wb.create_sheet("Scorecard")
    sc["A1"] = "Prediction scorecard"; sc["A1"].font = Font(name=F, bold=True, size=14)
    sc["A2"] = ("Every guess is logged before its draw and scored automatically when results arrive. "
                "Match columns are computed by the script on each run.")
    sc["A2"].font = Font(name=F, italic=True)
    heads = ["Game", "Draw date", "Strategy", "Guess (white balls)", "Bonus", "Actual white balls", "Actual bonus", "White matches", "Bonus match"]
    for c, h in enumerate(heads, start=1):
        cell = sc.cell(4, c, h); cell.font = hf; cell.fill = hfill; cell.alignment = ctr
    for i, s in enumerate(scored, start=5):
        sc.cell(i, 1, GAMES[s["game"]]["label"]); sc.cell(i, 2, s["target"]).number_format = "yyyy-mm-dd"
        sc.cell(i, 3, s["strategy"]); sc.cell(i, 4, " ".join(f"{x:02d}" for x in s["pw"])); sc.cell(i, 5, s["pb"])
        if s["actual"]:
            sc.cell(i, 6, " ".join(f"{x:02d}" for x in s["actual"][0])); sc.cell(i, 7, s["actual"][1])
            sc.cell(i, 8, s["wm"]); sc.cell(i, 9, "Yes" if s["bm"] else "No")
        else:
            sc.cell(i, 6, "awaiting draw")
        for c in range(1, 10):
            sc.cell(i, c).font = body
            if c != 3: sc.cell(i, c).alignment = Alignment(horizontal="center")
    last_log = 4 + len(scored)
    r0 = last_log + 3
    sc.cell(r0 - 1, 1, "Summary by strategy (scored guesses only)").font = Font(name=F, bold=True, size=12)
    for c, h in enumerate(["Game", "Strategy", "Guesses scored", "Avg white matches", "Random benchmark", "Bonus hits"], start=1):
        cell = sc.cell(r0, c, h); cell.font = hf; cell.fill = hfill; cell.alignment = ctr
    rr = r0 + 1
    rng = lambda col: f"${col}$5:${col}${max(last_log, 5)}"
    for g in preds:
        for name in ALL_NAMES:
            sc.cell(rr, 1, GAMES[g]["label"]); sc.cell(rr, 2, name)
            sc.cell(rr, 3, f'=COUNTIFS({rng("A")},A{rr},{rng("C")},B{rr},{rng("H")},">=0")')
            sc.cell(rr, 4, f'=IFERROR(AVERAGEIFS({rng("H")},{rng("A")},A{rr},{rng("C")},B{rr}),"")')
            sc.cell(rr, 5, f"=25/{GAMES[g]['whites']}")
            sc.cell(rr, 6, f'=COUNTIFS({rng("A")},A{rr},{rng("C")},B{rr},{rng("I")},"Yes")')
            sc.cell(rr, 4).number_format = "0.00"; sc.cell(rr, 5).number_format = "0.00"
            for c in range(1, 7):
                sc.cell(rr, c).font = body
                if c > 2: sc.cell(rr, c).alignment = Alignment(horizontal="center")
            rr += 1
    for col_, w in zip("ABCDEFGHI", [16, 34, 18, 24, 18, 24, 14, 14, 12]): sc.column_dimensions[col_].width = w
    sc.freeze_panes = "A5"

    # ---- About
    ab = wb.create_sheet("About")
    lines = ["About this workbook",
             "Rebuilt by lotto_predictor.py each run. Data source: NY Open Data lottery datasets (data.ny.gov).",
             f"Analysis window: last {days} days before the target draw. Mega Millions data starts 2025-04-08 (current number format).",
             "Strategies: " + "; ".join(f"{k} = {v}" for k, v in STRATEGY_HELP.items()),
             "Powerball and Mega Ball are drawn from separate drums, so the bonus ball can equal one of the white numbers.",
             "No strategy has predictive power. Odds: Powerball 1 in 292,201,338; Mega Millions 1 in 290,472,336."]
    for i, t in enumerate(lines, start=1):
        ab.cell(i, 1, t).font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 10)
        ab.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ab.column_dimensions["A"].width = 130
    wb.save(path)
    return True

def get_col(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)

# --------------------------------------------------------------------------- main
LEARN_FIELDS = ["run_at", "game", "target_date", "scope", "adopted", "wf", "wd", "wh", "half",
                "winners_rank_pct_default", "winners_rank_pct_chosen", "tests"]

def append_learning_log(path, game, target, tuned):
    new = not os.path.exists(path)
    with open(path, "a", newline="") as f:
        w = csv.DictWriter(f, fieldnames=LEARN_FIELDS)
        if new: w.writeheader()
        for scope in ("whites", "bonus"):
            t = tuned[scope]; wf, wd, wh, half = t["params"]
            w.writerow(dict(run_at=dt.datetime.now().isoformat(timespec="seconds"), game=game,
                            target_date=target.isoformat(), scope=scope, adopted=int(t["adopted"]),
                            wf=wf, wd=wd, wh=wh, half=half,
                            winners_rank_pct_default=round(t["default_metric"], 4),
                            winners_rank_pct_chosen=round(t["chosen_metric"], 4), tests=t["n"]))

def main(argv=None):
    ap = argparse.ArgumentParser(description="Lottery predictor (for fun)")
    ap.add_argument("--game", choices=["powerball", "megamillions", "both"], default="both")
    ap.add_argument("--days", type=int, default=365, help="history window in days (default 365)")
    ap.add_argument("--no-update", action="store_true", help="skip downloading new draws")
    ap.add_argument("--date", help="predict for the first draw on/after this date (YYYY-MM-DD); past dates allowed")
    ap.add_argument("--backtest", action="store_true", help="walk-forward test of every strategy")
    ap.add_argument("--data-dir", default=os.path.join(HERE, "data"))
    ap.add_argument("--out", default=os.path.join(HERE, "output"))
    ap.add_argument("--no-log", action="store_true", help="do not record these guesses in the prediction log")
    ap.add_argument("--no-learn", action="store_true", help="skip the self-tuning step (faster)")
    ap.add_argument("--learn-window", type=int, default=60, help="draws used to tune the adaptive strategy (default 60)")
    ap.add_argument("--review", nargs="?", const=3, type=int, metavar="K", help="show where the model went wrong on the last K draws (default 3)")
    ap.add_argument("--audit", nargs="?", const=12, type=int, metavar="N", help="out-of-sample test of self-tuning on the last N draws (default 12, slow)")
    ap.add_argument("--remind", action="store_true", help="only act if a draw is TOMORROW (ET); print a short reminder with the guesses")
    ap.add_argument("--ingest", nargs=2, action="append", metavar=("GAME", "FILE"),
                    help="add draws from a text/JSON file (use with --no-update when the feed is unreachable)")
    ap.add_argument("--no-workbook", action="store_true", help="do not write the Excel workbook")
    a = ap.parse_args(argv)
    games = list(GAMES) if a.game == "both" else [a.game]
    os.makedirs(a.out, exist_ok=True)

    all_draws, preds = {}, {}
    ingest = {g: f for g, f in (a.ingest or [])}
    for g in GAMES:
        if g in ingest:
            print(f"[{GAMES[g]['label']}] ingest: added {ingest_file(g, ingest[g], a.data_dir)} draw(s) from {ingest[g]}")
        if g in games and not a.no_update:
            n, msg = update_data(g, a.data_dir)
            print(f"[{GAMES[g]['label']}] update: {msg}")
        all_draws[g] = load_draws(g, a.data_dir)

    if a.backtest:
        for g in games:
            res, rw, rb = backtest(g, all_draws[g], a.days)
            cfg = GAMES[g]; n = len(rw)
            print(f"\n{cfg['label']} walk-forward backtest over {n} draws (uses only earlier draws):")
            print(f"  {'Strategy':34s} white/draw  bonus hit%  draws w/ 2+ whites")
            for name, _ in STRATEGIES:
                w, b = res[name]["w"], res[name]["b"]
                print(f"  {name:34s} {statistics.mean(w):8.3f}   {100 * sum(b) / n:7.1f}%   {sum(x >= 2 for x in w)}")
            print(f"  {'Random picks (sample)':34s} {statistics.mean(rw):8.3f}   {100 * sum(rb) / n:7.1f}%   {sum(x >= 2 for x in rw)}")
            print(f"  {'Theoretical random':34s} {25 / cfg['whites']:8.3f}   {100 / cfg['bonus']:7.1f}%")
        return 0

    if a.review is not None:
        for g in games: review(g, all_draws[g], a.days, a.review)
        return 0
    if a.audit is not None:
        for g in games: audit(g, all_draws[g], a.days, a.audit, a.learn_window)
        return 0

    start = dt.date.fromisoformat(a.date) if a.date else default_start()
    if a.remind:
        tomorrow = now_et().date() + dt.timedelta(days=1)
        start = tomorrow
        due = [g for g in games if tomorrow.weekday() in GAMES[g]["weekdays"]]
        if not due:
            nxt = ", ".join(f"{GAMES[g]['label']} {next_draw_date(g, tomorrow):%a %b %d}" for g in games)
            print(f"NO_DRAW_TOMORROW. No reminder needed. Next draws: {nxt}")
            return 0
        games = due
    log_path = os.path.join(a.data_dir, "picks_log.csv")
    log = load_log(log_path)
    logged = {(r["game"], r["target_date"], r["strategy"]) for r in log}
    new_log = []
    window_start = None
    for g in games:
        cfg = GAMES[g]
        target = next_draw_date(g, start)
        H, wstart = history_for(all_draws[g], target, a.days, cfg)
        if len(H) < 30:
            print(f"[{cfg['label']}] not enough history ({len(H)} draws) - need at least 30."); continue
        strategies = list(STRATEGIES); tuned = None; note = ""
        if not a.no_learn:
            tuned = tune(all_draws[g], cfg, target, a.days, a.learn_window)
            strategies.append((ADAPTIVE_NAME, make_adaptive(tuned)))
            note = learn_note(tuned)
        picks = []
        for name, f in strategies:
            combo, bonus = f(H, cfg); picks.append((name, combo, bonus))
        dates = [d for d, _, _ in all_draws[g] if wstart <= d < target]
        preds[g] = dict(target=target, picks=picks, n=len(H), start=min(dates), last=max(dates), learn_note=note)
        if tuned and not a.no_log:
            append_learning_log(os.path.join(a.data_dir, "learning_log.csv"), g, target, tuned)
        window_start = wstart if window_start is None else min(window_start, wstart)
        print(f"\n{cfg['label']} - draw {target:%a %Y-%m-%d} ({cfg['draw_time']}); {len(H)} draws from {min(dates)} to {max(dates)}")
        for name, combo, bonus in picks:
            print(f"  {name:34s} {' '.join(f'{x:02d}' for x in combo)}  + {bonus:02d}   (sum {sum(combo)})")
        if note: print(f"  Learning check: {note}")
        actual = next(((w, b) for d, w, b in all_draws[g] if d == target), None)
        if actual:
            print(f"  Actual result for {target}: {' '.join(f'{x:02d}' for x in actual[0])} + {actual[1]:02d}")
            for name, combo, bonus in picks:
                print(f"    {name:34s} white matches {len(set(combo) & set(actual[0]))}, {cfg['bonus_name']} {'hit' if bonus == actual[1] else 'miss'}")
        if not a.no_log and not actual:        # lock in guesses before the draw only
            for name, combo, bonus in picks:
                key = (g, target.isoformat(), name)
                if key not in logged:
                    new_log.append(dict(game=g, target_date=target.isoformat(), strategy=name,
                                        whites=" ".join(map(str, combo)), bonus=bonus,
                                        generated_at=dt.datetime.now().isoformat(timespec="seconds")))
    if new_log:
        append_log(log_path, new_log); print(f"\nLogged {len(new_log)} new guess(es) to {log_path}")

    scored = scored_log(load_log(log_path), all_draws)
    done = [s for s in scored if s["actual"]]
    if done:
        print("\nScorecard so far (logged guesses with results):")
        for g in games:
            for name in ALL_NAMES:
                rows = [s for s in done if s["game"] == g and s["strategy"] == name]
                if rows:
                    print(f"  {GAMES[g]['label']:14s} {name:34s} {len(rows)} scored, avg white matches "
                          f"{statistics.mean(s['wm'] for s in rows):.2f} (random {25 / GAMES[g]['whites']:.2f})")
    if a.remind and preds:
        lines = ["LOTTERY REMINDER (for-fun guesses; every combination has identical odds)"]
        for g, p in preds.items():
            cfg = GAMES[g]
            lines.append(f"{cfg['label']} draw TOMORROW, {p['target']:%a %b %d} at {cfg['draw_time']}. "
                         f"Based on {p['n']} draws ({p['start']} to {p['last']}).")
            for name, combo, bonus in p["picks"]:
                lines.append(f"  {name}: {' '.join(f'{x:02d}' for x in combo)} + {cfg['bonus_name']} {bonus:02d}")
            if p.get("learn_note"): lines.append("  Learning: " + p["learn_note"])
        msg = "\n".join(lines)
        print("\n" + msg)
        with open(os.path.join(a.out, "reminder.txt"), "w") as f: f.write(msg + "\n")
    if preds and not a.no_workbook:
        path = os.path.join(a.out, "lottery_workbook.xlsx")
        if write_workbook(path, {g: all_draws[g] for g in preds}, preds, [s for s in scored if s["game"] in preds],
                          a.days, window_start):
            print(f"\nWorkbook written: {path}")
    return 0

if __name__ == "__main__":
    sys.exit(main())
